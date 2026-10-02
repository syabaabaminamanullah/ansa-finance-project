from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

# Create Laporan directory if not exists
os.makedirs(r"D:\web dev - ansa\ansa-finance-project\Laporan", exist_ok=True)

doc = Document()

# Set font
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)

# --- HEADER (KOP SURAT) ---
section = doc.sections[0]
header = section.header
htable = header.add_table(1, 2, Inches(6))
htable.autofit = False
htable.columns[0].width = Inches(1.5)
htable.columns[1].width = Inches(4.5)

# Try to add logo
logo_path = r"D:\web dev - ansa\ansa-finance-project\frontend\public\logo-transparent.png"
if os.path.exists(logo_path):
    p_logo = htable.cell(0,0).paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_logo = p_logo.add_run()
    r_logo.add_picture(logo_path, width=Inches(1.2))

# Company details in header
p_hdr = htable.cell(0,1).paragraphs[0]
p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_hdr = p_hdr.add_run("PT CORETERRA GEO ENGINEERING\n")
run_hdr.bold = True
run_hdr.font.size = Pt(14)
p_hdr.add_run("General Contractor, Supplier, Geotechnical & Topography Survey\n")
p_hdr.add_run("Bandung, Jawa Barat")
# Add a bottom border to header
p_line = header.add_paragraph()
p_line.add_run("_" * 65).bold = True
p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER

# --- DOCUMENT BODY ---
# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('SURAT PERNYATAAN BUKAN PENGUSAHA KENA PAJAK (NON-PKP)')
run.bold = True
run.underline = True

doc.add_paragraph()

doc.add_paragraph("Yang bertanda tangan di bawah ini:")

# Add list
p = doc.add_paragraph()
p.add_run("• Nama\t\t\t: [Setyo Mardani / Nama Direktur]\n")
p.add_run("• Jabatan\t\t: Direktur\n")
p.add_run("• Nama Perusahaan\t: PT Coreterra Geo Engineering\n")
p.add_run("• Alamat Perusahaan\t: Bandung, Jawa Barat (Mohon lengkapi alamat detail)\n")
p.add_run("• NPWP Perusahaan\t: [Nomor NPWP Perusahaan]")

p_body = doc.add_paragraph()
p_body.add_run(
    "Dengan ini menyatakan bahwa sampai dengan saat ini, PT Coreterra Geo Engineering belum dikukuhkan "
    "sebagai Pengusaha Kena Pajak (PKP) oleh Direktorat Jenderal Pajak, dikarenakan jumlah peredaran bruto "
    "(omzet) perusahaan belum mencapai batasan wajib PKP sesuai peraturan perpajakan yang berlaku "
    "(di bawah Rp4,8 Miliar per tahun)."
)
p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

p_body2 = doc.add_paragraph("Oleh karena itu, dalam setiap transaksi penyerahan Barang/Jasa Kena Pajak, kami:")
p_body2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

doc.add_paragraph("1. Tidak memungut Pajak Pertambahan Nilai (PPN) 12%.")
doc.add_paragraph("2. Tidak menerbitkan Faktur Pajak.")

p_body3 = doc.add_paragraph(
    "Demikian surat pernyataan ini dibuat dengan sebenarnya untuk dapat dipergunakan sebagaimana mestinya. "
    "Jika di kemudian hari status perpajakan perusahaan kami berubah menjadi PKP, kami akan segera "
    "menginformasikan kepada pihak Mitra/Klien."
)
p_body3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

doc.add_paragraph()

# Signature section
p_sig = doc.add_paragraph(
    "Bandung, 1 Oktober 2026\n"
    "PT Coreterra Geo Engineering\n"
)
p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT

# Add Tanda Tangan
ttd_path = r"D:\web dev - ansa\ansa-finance-project\frontend\public\ttd-setyo.jpg"
cap_path = r"D:\web dev - ansa\ansa-finance-project\frontend\public\cap-transparent.png"

if os.path.exists(ttd_path):
    p_ttd = doc.add_paragraph()
    p_ttd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_ttd = p_ttd.add_run()
    r_ttd.add_picture(ttd_path, height=Inches(0.8))
else:
    doc.add_paragraph("\n\n(Meterai Rp10.000 & Tanda Tangan)\n\n").alignment = WD_ALIGN_PARAGRAPH.RIGHT

p_name = doc.add_paragraph("[Setyo Mardani / Nama Lengkap]\nDirektur")
p_name.alignment = WD_ALIGN_PARAGRAPH.RIGHT

doc.save(r"D:\web dev - ansa\ansa-finance-project\Laporan\Surat_Pernyataan_Non_PKP_Coreterra.docx")
print("Done")
