from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

# Create Laporan directory if not exists
os.makedirs(r"D:\web dev - ansa\ansa-finance-project\Laporan", exist_ok=True)

doc = Document()

# Set font
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)

# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('SURAT PERNYATAAN BUKAN PENGUSAHA KENA PAJAK (NON-PKP)')
run.bold = True
run.underline = True

doc.add_paragraph()

doc.add_paragraph("Yang bertanda tangan di bawah ini:")

# Add list
p = doc.add_paragraph()
p.add_run("• Nama\t\t\t: [Nama Direktur]\n")
p.add_run("• Jabatan\t\t: Direktur\n")
p.add_run("• Nama Perusahaan\t: PT Coreterra Geo Engineering\n")
p.add_run("• Alamat Perusahaan\t: [Alamat Lengkap Perusahaan]\n")
p.add_run("• NPWP Perusahaan\t: [Nomor NPWP Perusahaan]")

p_body = doc.add_paragraph()
p_body.add_run(
    "Dengan ini menyatakan bahwa sampai dengan saat ini, PT Coreterra Geo Engineering belum dikukuhkan "
    "sebagai Pengusaha Kena Pajak (PKP) oleh Direktorat Jenderal Pajak, dikarenakan jumlah peredaran bruto "
    "(omzet) perusahaan belum mencapai batasan wajib PKP sesuai peraturan perpajakan yang berlaku "
    "(di bawah Rp4,8 Miliar per tahun)."
)
p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

p_body2 = doc.add_paragraph("Oleh karena itu, dalam setiap transaksi penyerahan Barang/Jasa Kena Pajak, kami:")
p_body2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

doc.add_paragraph("1. Tidak memungut Pajak Pertambahan Nilai (PPN) 12%.")
doc.add_paragraph("2. Tidak menerbitkan Faktur Pajak.")

p_body3 = doc.add_paragraph(
    "Demikian surat pernyataan ini dibuat dengan sebenarnya untuk dapat dipergunakan sebagaimana mestinya. "
    "Jika di kemudian hari status perpajakan perusahaan kami berubah menjadi PKP, kami akan segera "
    "menginformasikan kepada pihak Mitra/Klien."
)
p_body3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

doc.add_paragraph()

# Signature section
p_sig = doc.add_paragraph("[Kota Perusahaan], 1 Oktober 2026\nPT Coreterra Geo Engineering\n\n\n\n\n(Meterai Rp10.000 & Tanda Tangan)\n\n[Nama Lengkap Direktur]\nDirektur")
p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT

doc.save(r"D:\web dev - ansa\ansa-finance-project\Laporan\Surat_Pernyataan_Non_PKP_Coreterra.docx")
print("Done")
